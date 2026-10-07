from __future__ import annotations

import asyncio
import hashlib
import io
import json
import math
import os
import re
import threading
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from pathlib import Path
from typing import Any

import pandas as pd
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy import Date, DateTime, Float, ForeignKey, Integer, String, Text, create_engine, func, select
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("DATA_DIR", ROOT / "data"))
INBOX = Path(os.getenv("INBOX_DIR", DATA_DIR / "inbox"))
REPORTS = Path(os.getenv("REPORT_DIR", DATA_DIR / "reports"))
for directory in (DATA_DIR, INBOX, REPORTS):
    directory.mkdir(parents=True, exist_ok=True)

DATABASE_URL = os.getenv("DATABASE_URL", f"sqlite:///{DATA_DIR / 'sales.db'}")
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {})
Session = sessionmaker(bind=engine, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class Run(Base):
    __tablename__ = "runs"
    id: Mapped[int] = mapped_column(primary_key=True)
    source: Mapped[str] = mapped_column(String(255))
    status: Mapped[str] = mapped_column(String(30), default="success")
    total_rows: Mapped[int] = mapped_column(Integer, default=0)
    accepted_rows: Mapped[int] = mapped_column(Integer, default=0)
    rejected_rows: Mapped[int] = mapped_column(Integer, default=0)
    duplicate_rows: Mapped[int] = mapped_column(Integer, default=0)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(ZoneInfo("America/Sao_Paulo")))
    errors: Mapped[list["ImportError"]] = relationship(back_populates="run", cascade="all, delete-orphan")


class Sale(Base):
    __tablename__ = "sales"
    id: Mapped[int] = mapped_column(primary_key=True)
    row_hash: Mapped[str] = mapped_column(String(64), unique=True, index=True)
    sale_date: Mapped[date] = mapped_column(Date)
    seller: Mapped[str] = mapped_column(String(160), index=True)
    product: Mapped[str] = mapped_column(String(200), index=True)
    quantity: Mapped[int] = mapped_column(Integer)
    unit_price: Mapped[float] = mapped_column(Float)
    source: Mapped[str] = mapped_column(String(255))
    run_id: Mapped[int] = mapped_column(ForeignKey("runs.id"))


class ImportError(Base):
    __tablename__ = "import_errors"
    id: Mapped[int] = mapped_column(primary_key=True)
    run_id: Mapped[int] = mapped_column(ForeignKey("runs.id"))
    row_number: Mapped[int] = mapped_column(Integer)
    reason: Mapped[str] = mapped_column(Text)
    raw_data: Mapped[str] = mapped_column(Text)
    run: Mapped[Run] = relationship(back_populates="errors")


Base.metadata.create_all(engine)
app = FastAPI(title="Relatório Automático de Vendas", version="1.0.0", description="Importe planilhas, valide vendas e gere indicadores automaticamente.")
app.mount("/assets", StaticFiles(directory=ROOT / "app" / "assets"), name="assets")

AUTO_PROCESS_INBOX = os.getenv("AUTO_PROCESS_INBOX", "true").strip().lower() not in {"0", "false", "no", "off"}
INBOX_POLL_SECONDS = max(2, int(os.getenv("INBOX_POLL_SECONDS", "5")))
INBOX_PROCESS_LOCK = threading.Lock()
AUTOMATION_STATE: dict[str, Any] = {"enabled": AUTO_PROCESS_INBOX, "interval_seconds": INBOX_POLL_SECONDS, "last_check": None, "last_event": None, "last_error": None, "event_id": 0}

ALIASES = {
    "data": {"data", "date", "data venda", "data da venda"},
    "vendedor": {"vendedor", "seller", "representante"},
    "produto": {"produto", "product", "item"},
    "quantidade": {"quantidade", "quantity", "qtd"},
    "valor_unitario": {"valor unitario", "preco unitario", "valor", "unit price", "price"},
}


def normalize(value: Any) -> str:
    raw = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", raw).strip()


def parse_price(value: Any) -> float:
    try:
        if isinstance(value, (int, float, Decimal)):
            amount = Decimal(str(value))
        else:
            text = str(value).strip().replace("R$", "").replace("\u00a0", "").replace(" ", "")
            if "," in text and "." in text:
                decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
                grouping_separator = "." if decimal_separator == "," else ","
                text = text.replace(grouping_separator, "").replace(decimal_separator, ".")
            elif "," in text:
                text = text.replace(".", "").replace(",", ".")
            amount = Decimal(text)
        if not amount.is_finite():
            raise ValueError("Invalid currency amount")
        return float(amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
    except InvalidOperation as exc:
        raise ValueError("Invalid currency amount") from exc


def money_decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def read_frame(content: bytes, filename: str) -> pd.DataFrame:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            try:
                return pd.read_csv(io.BytesIO(content), sep=None, engine="python")
            except UnicodeDecodeError:
                return pd.read_csv(io.BytesIO(content), sep=None, engine="python", encoding="latin-1")
        if suffix in {".xlsx", ".xlsm"}:
            return pd.read_excel(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(422, f"Não foi possível ler a planilha: {exc}") from exc
    raise HTTPException(400, "Formato não suportado. Envie .xlsx, .xlsm ou .csv.")


def process_bytes(content: bytes, filename: str) -> dict[str, Any]:
    frame = read_frame(content, filename)
    source = Path(filename).name[:255]
    columns = {normalize(column): column for column in frame.columns}
    mapped: dict[str, Any] = {}
    for target, aliases in ALIASES.items():
        match = next((columns[name] for name in aliases if name in columns), None)
        if match is None:
            raise HTTPException(422, f"Coluna obrigatória ausente: {target}. Colunas aceitas: data, vendedor, produto, quantidade, valor_unitario.")
        mapped[target] = match

    with Session() as db:
        run = Run(source=source, status="processing", total_rows=len(frame))
        db.add(run)
        db.flush()
        seen_in_file: set[str] = set()
        for index, row in frame.iterrows():
            raw = {str(k): (None if pd.isna(v) else str(v)) for k, v in row.items()}
            try:
                sale_date = pd.to_datetime(row[mapped["data"]], dayfirst=True, errors="raise").date()
                seller, product = str(row[mapped["vendedor"]]).strip(), str(row[mapped["produto"]]).strip()
                quantity_value = float(row[mapped["quantidade"]])
                if not quantity_value.is_integer():
                    raise ValueError("quantidade precisa ser um número inteiro")
                quantity = int(quantity_value)
                price = parse_price(row[mapped["valor_unitario"]])
                if not seller or seller.lower() == "nan":
                    raise ValueError("vendedor vazio")
                if not product or product.lower() == "nan":
                    raise ValueError("produto vazio")
                if quantity <= 0:
                    raise ValueError("quantidade deve ser maior que zero")
                if price < 0:
                    raise ValueError("valor unitário não pode ser negativo")
                canonical = f"{sale_date.isoformat()}|{seller.casefold()}|{product.casefold()}|{quantity}|{price:.4f}"
                digest = hashlib.sha256(canonical.encode()).hexdigest()
                exists = digest in seen_in_file or db.scalar(select(Sale.id).where(Sale.row_hash == digest)) is not None
                if exists:
                    run.duplicate_rows += 1
                    continue
                seen_in_file.add(digest)
                db.add(Sale(row_hash=digest, sale_date=sale_date, seller=seller, product=product, quantity=quantity, unit_price=price, source=source, run_id=run.id))
                run.accepted_rows += 1
            except Exception as exc:
                run.rejected_rows += 1
                db.add(ImportError(run_id=run.id, row_number=int(index) + 2, reason=str(exc), raw_data=json.dumps(raw, ensure_ascii=False)))
        run.status = "warning" if run.rejected_rows or run.duplicate_rows else "success"
        db.commit()
        result = run_dict(run)

    export_report()
    source_dir = DATA_DIR / "sources"
    source_dir.mkdir(parents=True, exist_ok=True)
    source_suffix = Path(filename).suffix.lower() if Path(filename).suffix.lower() in {".xlsx", ".xlsm", ".csv"} else ".xlsx"
    (source_dir / f"{result['id']}{source_suffix}").write_bytes(content)
    return result


def run_timestamp_local(value: datetime | None) -> str | None:
    if value is None:
        return None
    local_zone = ZoneInfo("America/Sao_Paulo")
    if value.tzinfo is None:
        value = value.replace(tzinfo=local_zone)
    else:
        value = value.astimezone(local_zone)
    return value.isoformat()


def run_dict(run: Run) -> dict[str, Any]:
    return {"id": run.id, "source": run.source, "status": run.status, "total_rows": run.total_rows,
            "accepted_rows": run.accepted_rows, "rejected_rows": run.rejected_rows,
            "duplicate_rows": run.duplicate_rows, "created_at": run_timestamp_local(run.created_at),
            "errors": [{"row_number": e.row_number, "reason": e.reason, "raw_data": json.loads(e.raw_data)} for e in run.errors]}


def summary() -> dict[str, Any]:
    with Session() as db:
        count, units, revenue = db.execute(select(func.count(Sale.id), func.coalesce(func.sum(Sale.quantity), 0), func.coalesce(func.sum(Sale.quantity * Sale.unit_price), 0))).one()
        by_seller = db.execute(select(Sale.seller, func.count(Sale.id), func.sum(Sale.quantity * Sale.unit_price)).group_by(Sale.seller).order_by(func.sum(Sale.quantity * Sale.unit_price).desc())).all()
        by_product = db.execute(select(Sale.product, func.sum(Sale.quantity), func.sum(Sale.quantity * Sale.unit_price)).group_by(Sale.product).order_by(func.sum(Sale.quantity).desc()).limit(10)).all()
        latest_date = db.scalar(select(func.max(Sale.sale_date)))
        trend = []
        if latest_date:
            first_date = latest_date - timedelta(days=119)
            daily = db.execute(select(Sale.sale_date, func.sum(Sale.quantity * Sale.unit_price)).where(Sale.sale_date >= first_date, Sale.sale_date <= latest_date).group_by(Sale.sale_date).order_by(Sale.sale_date)).all()
            trend = [{"date": day.isoformat(), "revenue": round(float(value or 0), 2)} for day, value in daily]
        return {"sale_count": count, "units_sold": int(units), "revenue": float(money_decimal(revenue)), "average_sale": float(money_decimal(Decimal(str(revenue or 0)) / count)) if count else 0,
                "by_seller": [{"seller": s, "sales": int(n), "revenue": float(money_decimal(r))} for s, n, r in by_seller],
                "top_products": [{"product": p, "units": int(q or 0), "revenue": float(money_decimal(r))} for p, q, r in by_product],
                "trend": trend, "latest_sale_date": latest_date.isoformat() if latest_date else None}


def export_report() -> None:
    data = summary()
    data["generated_at"] = datetime.now(timezone.utc).isoformat()
    (REPORTS / "relatorio.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    pd.DataFrame(data["by_seller"]).to_csv(REPORTS / "vendas_por_vendedor.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(data["top_products"]).to_csv(REPORTS / "produtos_mais_vendidos.csv", index=False, encoding="utf-8-sig")


@app.get("/", response_class=HTMLResponse)
def dashboard():
    return FileResponse(ROOT / "app" / "index.html")


@app.get("/api/health")
def health():
    return {"status": "ok", "database": "connected"}


@app.get("/api/automation/status")
def automation_status():
    return {key: value for key, value in AUTOMATION_STATE.items() if key != "event_id"}


@app.post("/api/process/upload")
async def upload(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(413, "O arquivo excede o limite de 25 MB.")
    return process_bytes(content, file.filename or "upload.xlsx")


def process_inbox_files(paths: list[Path] | None = None, trigger: str = "manual") -> dict[str, Any]:
    if not INBOX_PROCESS_LOCK.acquire(blocking=False):
        return {"files_found": 0, "runs": [], "busy": True}
    results = []
    try:
        candidates = sorted(paths if paths is not None else INBOX.iterdir())
        for path in candidates:
            if not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
                continue
            content = path.read_bytes()
            try:
                result = process_bytes(content, path.name)
                archive_dir = INBOX / "processed"
            except Exception as exc:
                detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
                message = f"Não foi possível processar o arquivo. Confira o formato e as colunas obrigatórias. ({detail})"
                with Session() as db:
                    run = Run(source=path.name[:255], status="error", total_rows=0, rejected_rows=1)
                    db.add(run)
                    db.flush()
                    db.add(ImportError(run_id=run.id, row_number=0, reason=message, raw_data=json.dumps({"arquivo": path.name}, ensure_ascii=False)))
                    db.commit()
                    result = run_dict(run)
                source_dir = DATA_DIR / "sources"
                source_dir.mkdir(parents=True, exist_ok=True)
                (source_dir / f"{result['id']}{path.suffix.lower()}").write_bytes(content)
                archive_dir = INBOX / "failed"
            archive_dir.mkdir(exist_ok=True)
            destination = archive_dir / path.name
            if destination.exists():
                destination = archive_dir / f"{path.stem}_{datetime.now().strftime('%Y%m%d%H%M%S')}{path.suffix}"
            path.replace(destination)
            results.append(result)

        if trigger == "automatic" and results:
            AUTOMATION_STATE["event_id"] += 1
            AUTOMATION_STATE["last_event"] = {
                "id": datetime.now(timezone.utc).isoformat(),
                "created_at": datetime.now(timezone.utc).isoformat(),
                "files": [run.get("source", "arquivo") for run in results],
                "files_processed": len(results),
                "accepted_rows": sum(run.get("accepted_rows", 0) for run in results),
                "rejected_rows": sum(run.get("rejected_rows", 0) for run in results),
                "duplicate_rows": sum(run.get("duplicate_rows", 0) for run in results),
                "has_errors": any(run.get("status") == "error" for run in results),
            }
            AUTOMATION_STATE["last_error"] = None
        return {"files_found": len(results), "runs": results, "busy": False}
    finally:
        INBOX_PROCESS_LOCK.release()


async def monitor_inbox() -> None:
    stable_files: dict[str, tuple[int, int]] = {}
    while True:
        try:
            ready = []
            current_files = {}
            for path in INBOX.iterdir():
                if not path.is_file() or path.name.startswith(("~$", ".")) or path.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
                    continue
                stat = path.stat()
                signature = (stat.st_size, stat.st_mtime_ns)
                current_files[str(path)] = signature
                if stable_files.get(str(path)) == signature:
                    ready.append(path)
            stable_files = current_files
            AUTOMATION_STATE["last_check"] = datetime.now(timezone.utc).isoformat()
            if ready:
                await asyncio.to_thread(process_inbox_files, ready, "automatic")
            AUTOMATION_STATE["last_error"] = None
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            AUTOMATION_STATE["last_error"] = str(exc)
        await asyncio.sleep(INBOX_POLL_SECONDS)


@app.on_event("startup")
async def start_inbox_monitor() -> None:
    if AUTO_PROCESS_INBOX:
        app.state.inbox_monitor_task = asyncio.create_task(monitor_inbox())


@app.on_event("shutdown")
async def stop_inbox_monitor() -> None:
    task = getattr(app.state, "inbox_monitor_task", None)
    if task:
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass


@app.post("/api/process/inbox")
def process_inbox():
    return process_inbox_files(trigger="manual")


@app.get("/api/reports/summary")
def get_summary():
    return summary()


@app.get("/api/runs")
def list_runs():
    with Session() as db:
        runs = db.scalars(select(Run).order_by(Run.created_at.desc()).limit(100)).all()
        return [run_dict(run) for run in runs]


def find_source_file(run: Run) -> Path | None:
    source_dir = DATA_DIR / "sources"
    for suffix in (".xlsx", ".xlsm", ".csv"):
        candidate = source_dir / f"{run.id}{suffix}"
        if candidate.is_file():
            return candidate
    for directory in (INBOX / "processed", INBOX / "failed", ROOT / "examples"):
        candidate = directory / run.source
        if candidate.is_file():
            return candidate
    return None


def row_values(raw_data: dict[str, Any]) -> dict[str, Any]:
    values = {}
    for target, aliases in ALIASES.items():
        key = next((name for name in raw_data if normalize(name) in aliases), None)
        values[target] = raw_data.get(key) if key is not None else None
    return values


def parse_corrections(run: Run, payload: dict[str, Any]) -> list[dict[str, Any]]:
    errors = [error for error in run.errors if error.row_number > 0]
    submitted = {int(item.get("row_number", -1)): item for item in payload.get("corrections", []) if isinstance(item, dict)}
    corrections = []
    for error in errors:
        item = submitted.get(error.row_number)
        if not item:
            raise HTTPException(422, f"Preencha os campos da linha {error.row_number} antes de exportar.")
        try:
            sale_date = date.fromisoformat(str(item.get("data", "")).strip())
            seller = str(item.get("vendedor", "")).strip()
            product = str(item.get("produto", "")).strip()
            quantity_number = float(str(item.get("quantidade", "")).strip().replace(",", "."))
            price = parse_price(item.get("valor_unitario", ""))
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, f"Revise os campos da linha {error.row_number}; confira data, quantidade e valor unitário.") from exc
        if not seller or not product:
            raise HTTPException(422, f"Informe vendedor e produto na linha {error.row_number}.")
        if not math.isfinite(quantity_number) or not quantity_number.is_integer() or quantity_number <= 0:
            raise HTTPException(422, f"A quantidade da linha {error.row_number} deve ser um número inteiro maior que zero.")
        if not math.isfinite(price) or price < 0:
            raise HTTPException(422, f"O valor unitário da linha {error.row_number} deve ser zero ou maior.")
        corrections.append({"row_number": error.row_number, "data": sale_date, "vendedor": seller, "produto": product, "quantidade": int(quantity_number), "valor_unitario": float(price)})
    return corrections


def corrections_for_export(run: Run, payload: dict[str, Any]) -> list[dict[str, Any]]:
    if not run.rejected_rows:
        return []
    if not run.errors or any(error.row_number <= 0 for error in run.errors):
        raise HTTPException(422, "Esta planilha tem um erro de estrutura e não pode ser exportada corrigida. Ajuste os cabeçalhos ou o formato e importe novamente.")
    return parse_corrections(run, payload)


def run_sales_rows(run_id: int, corrections: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    with Session() as db:
        sales = db.scalars(select(Sale).where(Sale.run_id == run_id).order_by(Sale.sale_date, Sale.id)).all()
        rows = [{"data": sale.sale_date, "vendedor": sale.seller, "produto": sale.product, "quantidade": sale.quantity, "valor_unitario": sale.unit_price} for sale in sales]
    known = {f"{row['data'].isoformat()}|{row['vendedor'].casefold()}|{row['produto'].casefold()}|{row['quantidade']}|{row['valor_unitario']:.4f}" for row in rows}
    for correction in corrections or []:
        canonical = f"{correction['data'].isoformat()}|{correction['vendedor'].casefold()}|{correction['produto'].casefold()}|{correction['quantidade']}|{correction['valor_unitario']:.4f}"
        if canonical in known:
            continue
        known.add(canonical)
        rows.append(correction)
    return sorted(rows, key=lambda row: (row["data"], row["vendedor"].casefold(), row["produto"].casefold()))


def dashboard_rows_for_run(run: Run, corrections: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int, int]:
    source = find_source_file(run)
    if not source:
        rows = run_sales_rows(run.id, corrections)
        with Session() as db:
            omitted_count = db.scalar(select(func.count(ImportError.id)).where(ImportError.run_id == run.id, ImportError.row_number > 0)) or 0
            stored_run = db.get(Run, run.id)
            omitted_count += 1 if stored_run.status == "error" else 0
            duplicate_count = stored_run.duplicate_rows
        return rows, max(omitted_count - len(corrections), 0), duplicate_count

    frame = read_frame(source.read_bytes(), source.name)
    available = {normalize(column): column for column in frame.columns}
    columns = {target: next((available[name] for name in aliases if name in available), None) for target, aliases in ALIASES.items()}
    if any(column is None for column in columns.values()):
        rows = run_sales_rows(run.id, corrections)
        with Session() as db:
            omitted_count = db.scalar(select(func.count(ImportError.id)).where(ImportError.run_id == run.id, ImportError.row_number > 0)) or 0
            stored_run = db.get(Run, run.id)
            omitted_count += 1 if stored_run.status == "error" else 0
            duplicate_count = stored_run.duplicate_rows
        return rows, max(omitted_count - len(corrections), 0), duplicate_count

    by_row = {item["row_number"]: item for item in corrections}
    rows, seen, omitted, duplicates = [], set(), 0, 0
    for index, (_, source_row) in enumerate(frame.iterrows(), start=2):
        correction = by_row.get(index)
        try:
            if correction:
                row = correction
            else:
                raw_date = source_row[columns["data"]]
                seller = "" if pd.isna(source_row[columns["vendedor"]]) else str(source_row[columns["vendedor"]]).strip()
                product = "" if pd.isna(source_row[columns["produto"]]) else str(source_row[columns["produto"]]).strip()
                quantity_value = float(str(source_row[columns["quantidade"]]).replace(",", "."))
                price = parse_price(source_row[columns["valor_unitario"]])
                sale_date = pd.to_datetime(raw_date, dayfirst=True, errors="raise").date()
                if not seller or not product or not math.isfinite(quantity_value) or not quantity_value.is_integer() or quantity_value <= 0 or not math.isfinite(price) or price < 0:
                    raise ValueError("Invalid sale row")
                row = {"data": sale_date, "vendedor": seller, "produto": product, "quantidade": int(quantity_value), "valor_unitario": float(price)}
            canonical = f"{row['data'].isoformat()}|{row['vendedor'].casefold()}|{row['produto'].casefold()}|{row['quantidade']}|{row['valor_unitario']:.4f}"
            if canonical in seen:
                duplicates += 1
                continue
            seen.add(canonical)
            rows.append(row)
        except Exception:
            omitted += 1
    return rows, omitted, max(duplicates, run.duplicate_rows)


def workbook_column_map(headers: list[Any]) -> dict[str, int]:
    normalized = {normalize(value): index + 1 for index, value in enumerate(headers)}
    mapping = {}
    for target, aliases in ALIASES.items():
        match = next((normalized[name] for name in aliases if name in normalized), None)
        if match:
            mapping[target] = match
    return mapping


def corrected_workbook(run: Run, corrections: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    source = find_source_file(run)
    output = io.BytesIO()
    if source and source.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(source)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        columns = workbook_column_map(headers)
        if len(columns) < len(ALIASES):
            raise HTTPException(422, "Não foi possível localizar todos os cabeçalhos para corrigir esta planilha.")
        for item in corrections:
            values = {"data": item["data"], "vendedor": item["vendedor"], "produto": item["produto"], "quantidade": item["quantidade"], "valor_unitario": item["valor_unitario"]}
            for target, value in values.items():
                cell = sheet.cell(row=item["row_number"], column=columns[target])
                cell.value = value
                if target == "data":
                    cell.number_format = "dd/mm/yyyy"
        for row_number in range(2, sheet.max_row + 1):
            price_cell = sheet.cell(row=row_number, column=columns["valor_unitario"])
            if price_cell.value is not None:
                try:
                    price_cell.value = parse_price(price_cell.value)
                except ValueError:
                    pass
                price_cell.number_format = 'R$ #,##0.00'
        workbook.save(output)
        return output.getvalue()

    if source and source.suffix.lower() == ".csv":
        frame = read_frame(source.read_bytes(), source.name)
        headers = list(frame.columns)
        columns = workbook_column_map(headers)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Vendas"
        sheet.append(headers)
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in values])
        for item in corrections:
            for target, value in {"data": item["data"], "vendedor": item["vendedor"], "produto": item["produto"], "quantidade": item["quantidade"], "valor_unitario": item["valor_unitario"]}.items():
                cell = sheet.cell(row=item["row_number"], column=columns[target])
                cell.value = value
                if target == "data":
                    cell.number_format = "dd/mm/yyyy"
        for row_number in range(2, sheet.max_row + 1):
            price_cell = sheet.cell(row=row_number, column=columns["valor_unitario"])
            if price_cell.value is not None:
                try:
                    price_cell.value = parse_price(price_cell.value)
                except ValueError:
                    pass
                price_cell.number_format = 'R$ #,##0.00'
    else:
        rows = run_sales_rows(run.id, corrections)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Vendas"
        headers = ["data", "vendedor", "produto", "quantidade", "valor_unitario"]
        sheet.append(headers)
        for row in rows:
            sheet.append([row["data"], row["vendedor"], row["produto"], row["quantidade"], row["valor_unitario"]])
            sheet.cell(row=sheet.max_row, column=1).number_format = "dd/mm/yyyy"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17466D")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 32)
    sheet.freeze_panes = "A2"
    workbook.save(output)
    return output.getvalue()


def dashboard_html(rows: list[dict[str, Any]], omitted_rows: int, duplicates: int) -> str:
    import html

    def brl(value: float) -> str:
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    count = len(rows)
    units = sum(int(row["quantidade"]) for row in rows)
    revenue = sum((Decimal(int(row["quantidade"])) * money_decimal(row["valor_unitario"]) for row in rows), Decimal("0.00"))
    average = money_decimal(revenue / count) if count else Decimal("0.00")
    sellers: dict[str, dict[str, Any]] = {}
    products: dict[str, dict[str, Any]] = {}
    days: dict[str, Decimal] = {}
    for row in rows:
        sale_total = Decimal(int(row["quantidade"])) * money_decimal(row["valor_unitario"])
        seller = sellers.setdefault(row["vendedor"], {"sales": 0, "revenue": Decimal("0.00")})
        seller["sales"] += 1
        seller["revenue"] += sale_total
        product = products.setdefault(row["produto"], {"units": 0, "revenue": Decimal("0.00")})
        product["units"] += int(row["quantidade"])
        product["revenue"] += sale_total
        key = row["data"].isoformat()
        days[key] = days.get(key, Decimal("0.00")) + sale_total
    seller_rows = "".join(f"<tr><td>{html.escape(name)}</td><td>{item['sales']}</td><td>{brl(item['revenue'])}</td></tr>" for name, item in sorted(sellers.items(), key=lambda item: item[1]["revenue"], reverse=True))
    product_rows = "".join(f"<tr><td>{html.escape(name)}</td><td>{item['units']}</td><td>{brl(item['revenue'])}</td></tr>" for name, item in sorted(products.items(), key=lambda item: item[1]["units"], reverse=True)[:10])
    points = sorted(days.items())[-30:]
    if points:
        max_day = max(value for _, value in points) or 1
        coords = [(36 + index * (728 / max(len(points) - 1, 1)), 182 - value / max_day * 148) for index, (_, value) in enumerate(points)]
        path = " ".join(("M" if index == 0 else "L") + f"{x:.1f},{y:.1f}" for index, (x, y) in enumerate(coords))
        chart = f'<svg viewBox="0 0 800 220" role="img" aria-label="Faturamento por dia"><path d="{path} L{coords[-1][0]:.1f},190 L{coords[0][0]:.1f},190 Z" fill="#2563eb" opacity=".10"/><path d="{path}" fill="none" stroke="#2563eb" stroke-width="4" stroke-linecap="round" stroke-linejoin="round"/>' + "".join(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="4" fill="#2563eb"><title>{html.escape(day)}: R$ {brl(value)}</title></circle>' for (day, value), (x, y) in zip(points, coords)) + '</svg>'
    else:
        chart = '<p class="empty">Não há vendas válidas para desenhar o gráfico.</p>'
    note = f"{omitted_rows} linha(s) rejeitada(s) não incluídas." if omitted_rows else "Todos os registros válidos desta importação estão no painel."
    if duplicates:
        note += f" {duplicates} duplicata(s) ignorada(s)."
    escaped_note = html.escape(note)
    return f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Dashboard de vendas</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f6fb;color:#1b2b43;font:15px Inter,Segoe UI,Arial,sans-serif}}main{{max-width:1400px;margin:auto;padding:38px}}header{{display:flex;justify-content:space-between;align-items:end;gap:20px;margin-bottom:26px}}.eyebrow{{color:#2563eb;font-size:11px;font-weight:800;letter-spacing:.13em;text-transform:uppercase}}h1{{margin:8px 0;font-size:34px;letter-spacing:-.04em}}p{{color:#74839a}}.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:15px}}.card,.panel{{background:#fff;border:1px solid #e1e8f1;border-radius:17px;box-shadow:0 10px 26px #1733540c}}.card{{padding:21px;background:radial-gradient(ellipse at 100% 0%,#2563eb14,transparent 55%),#fff}}.card span{{color:#74839a;font-size:12px}}.card strong{{display:block;margin-top:13px;font-size:27px}}.panel{{margin-top:18px;padding:22px}}.chart{{height:250px}}.chart svg{{width:100%;height:100%}}.tables{{display:grid;grid-template-columns:1fr 1fr;gap:18px}}h2{{font-size:17px;margin:0 0 16px}}table{{width:100%;border-collapse:collapse}}th,td{{padding:11px;text-align:left;border-bottom:1px solid #edf0f5}}th{{font-size:11px;text-transform:uppercase;color:#78879b}}td:last-child{{text-align:right}}.note{{margin-top:18px;padding:12px 14px;border-radius:10px;background:#eef3fb;color:#61738d;font-size:12px}}.empty{{color:#8290a3;padding:20px}}@media(max-width:850px){{main{{padding:24px}}.cards{{grid-template-columns:repeat(2,1fr)}}.tables{{grid-template-columns:1fr}}}}@media(max-width:520px){{main{{padding:16px}}header{{align-items:start;flex-direction:column}}h1{{font-size:28px}}.cards{{gap:9px}}.card{{padding:15px}}.card strong{{font-size:21px}}.panel{{padding:15px}}}}
</style></head><body><main><header><div><div class="eyebrow">SalesFlow · relatório exportado</div><h1>Dashboard de vendas</h1><p>Resumo gerado a partir desta importação.</p></div><div>{count} vendas válidas</div></header><section class="cards"><article class="card"><span>Faturamento</span><strong>R$ {brl(revenue)}</strong></article><article class="card"><span>Vendas</span><strong>{count}</strong></article><article class="card"><span>Unidades</span><strong>{units}</strong></article><article class="card"><span>Ticket médio</span><strong>R$ {brl(average)}</strong></article></section><section class="panel"><h2>Faturamento por dia</h2><div class="chart">{chart}</div></section><section class="tables"><article class="panel"><h2>Vendas por vendedor</h2><table><thead><tr><th>Vendedor</th><th>Vendas</th><th>Faturamento (R$)</th></tr></thead><tbody>{seller_rows or '<tr><td colspan="3">Sem dados</td></tr>'}</tbody></table></article><article class="panel"><h2>Produtos mais vendidos</h2><table><thead><tr><th>Produto</th><th>Unidades</th><th>Faturamento (R$)</th></tr></thead><tbody>{product_rows or '<tr><td colspan="3">Sem dados</td></tr>'}</tbody></table></article></section><div class="note">{escaped_note}</div></main></body></html>'''


@app.post("/api/runs/{run_id}/corrected-export")
def export_corrected_run(run_id: int, payload: dict[str, Any]):
    with Session() as db:
        run = db.get(Run, run_id)
        if not run:
            raise HTTPException(404, "Execução não encontrada.")
        corrections = corrections_for_export(run, payload)
        content = corrected_workbook(run, corrections)
    filename = f"vendas_corrigidas_{run_id}.xlsx"
    return StreamingResponse(io.BytesIO(content), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.post("/api/runs/{run_id}/dashboard-export")
def export_run_dashboard(run_id: int, payload: dict[str, Any] | None = None):
    raise HTTPException(410, "HTML dashboard export is disabled. Export the Power BI project instead.")
    payload = payload or {}
    with Session() as db:
        run = db.get(Run, run_id)
        if not run:
            raise HTTPException(404, "Execução não encontrada.")
        corrections = corrections_for_export(run, payload)
    rows, omitted_rows, duplicate_rows = dashboard_rows_for_run(run, corrections)
    content = dashboard_html(rows, max(omitted_rows, 0), duplicate_rows)
    filename = f"dashboard_vendas_{run_id}.html"
    return HTMLResponse(content, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def powerbi_project(rows: list[dict[str, Any]], run_id: int) -> bytes:
    """Build a self-contained PBIP project that Power BI Desktop can open and save as PBIX."""
    project = "SalesFlow_Vendas"
    model_name = f"{project}.SemanticModel"
    report_name = f"{project}.Report"
    sales_rows = []
    for row in rows:
        d = row["data"]
        if isinstance(d, str):
            d = date.fromisoformat(d)
        sales_rows.append((d, str(row["vendedor"]), str(row["produto"]), int(row["quantidade"]), float(row["valor_unitario"])))
    def dax_string(value: str) -> str:
        return '"' + value.replace('"', '""') + '"'
    dax_rows = ",\n        ".join(
        f'{{dt"{d.isoformat()}", {dax_string(seller)}, {dax_string(product)}, {quantity}, {price:.4f}}}'
        for d, seller, product, quantity, price in sales_rows
    ) or '{dt"2000-01-01", "Sem dados", "Sem dados", 0, 0}'
    dax_expression = (
        "DATATABLE(\"Data\", DATETIME, \"Vendedor\", STRING, \"Produto\", STRING, "
        "\"Quantidade\", INTEGER, \"Valor unit\u00e1rio\", DOUBLE, {\n        "
        + dax_rows
        + "\n    })"
    )
    model = {
        "name": project,
        "compatibilityLevel": 1600,
        "model": {
            "culture": "pt-BR",
            "tables": [{
                "name": "Vendas",
                "columns": [
                    {"name": "Data", "dataType": "dateTime", "sourceColumn": "Data", "formatString": "dd/mm/yyyy"},
                    {"name": "Vendedor", "dataType": "string", "sourceColumn": "Vendedor"},
                    {"name": "Produto", "dataType": "string", "sourceColumn": "Produto"},
                    {"name": "Quantidade", "dataType": "int64", "sourceColumn": "Quantidade", "formatString": "#,0"},
                    {"name": "Valor unitário", "dataType": "double", "sourceColumn": "Valor unitário", "formatString": "R$ #,##0.00"},
                ],
                "measures": [
                    {"name": "Faturamento", "expression": "SUMX(Vendas, Vendas[Quantidade] * Vendas[Valor unitário])", "formatString": "R$ #,##0.00"},
                    {"name": "Vendas", "expression": "COUNTROWS(Vendas)", "formatString": "#,##0"},
                    {"name": "Unidades", "expression": "SUM(Vendas[Quantidade])", "formatString": "#,##0"},
                    {"name": "Ticket médio", "expression": "DIVIDE([Faturamento], [Vendas], 0)", "formatString": "R$ #,##0.00"},
                ],
                "partitions": [{"name": "Vendas", "mode": "import", "source": {"type": "calculated", "expression": dax_expression}}],
            }],
        },
    }
    model_json = json.dumps(model, ensure_ascii=False, indent=2)
    pbip = {"$schema": "https://developer.microsoft.com/json-schemas/fabric/pbip/pbipProperties/1.0.0/schema.json", "version": "1.0", "artifacts": [{"report": {"path": report_name}}], "settings": {"enableAutoRecovery": True}}
    pbir = {"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definitionProperties/2.0.0/schema.json", "version": "4.0", "datasetReference": {"byPath": {"path": f"../{model_name}"}}}
    def source_col(column: str) -> dict[str, Any]:
        return {"Column": {"Expression": {"SourceRef": {"Entity": "Vendas"}}, "Property": column}}
    def measure_field(measure: str) -> dict[str, Any]:
        return {"Measure": {"Expression": {"SourceRef": {"Entity": "Vendas"}}, "Property": measure}}
    def projection(field: dict[str, Any], query_ref: str) -> dict[str, Any]:
        return {"field": field, "queryRef": query_ref, "active": True}
    def literal(value: str) -> dict[str, Any]:
        return {"expr": {"Literal": {"Value": value}}}
    def text_literal(value: str) -> dict[str, Any]:
        return literal("'" + value.replace("'", "''") + "'")
    def fill(color: str) -> dict[str, Any]:
        return {"solid": {"color": text_literal(color)}}
    def visual(name: str, visual_type: str, x: int, y: int, width: int, height: int, roles: dict[str, list[dict[str, Any]]], title: str) -> dict[str, Any]:
        return {"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/visualContainer/2.7.0/schema.json", "name": name, "position": {"x": x, "y": y, "z": 1000, "width": width, "height": height, "tabOrder": y * 10 + x}, "visual": {"visualType": visual_type, "query": {"queryState": {role: {"projections": fields} for role, fields in roles.items()}}, "drillFilterOtherVisuals": True, "visualContainerObjects": {"title": [{"properties": {"show": literal("false" if visual_type == "card" else "true"), "text": text_literal(title), "fontColor": fill("#172B46"), "fontSize": literal("13D"), "fontFamily": text_literal("Segoe UI Semibold"), "bold": literal("true"), "alignment": text_literal("left")}}], "background": [{"properties": {"show": literal("true"), "color": fill("#FFFFFF"), "transparency": literal("0D")}}], "border": [{"properties": {"show": literal("true"), "color": fill("#E4EAF1"), "width": literal("1D"), "radius": literal("14D")}}], "dropShadow": [{"properties": {"show": literal("true"), "color": fill("#172B46"), "transparency": literal("90D"), "position": text_literal("Outer"), "preset": text_literal("BottomRight")}}]}}}
    def field_measure(name: str) -> dict[str, Any]:
        return projection(measure_field(name), f"Vendas.{name}")
    def field_column(name: str) -> dict[str, Any]:
        return projection(source_col(name), f"Vendas.{name}")
    cards = [visual("Card_Faturamento", "card", 28, 22, 290, 125, {"Fields": [field_measure("Faturamento")]}, "Faturamento"), visual("Card_Vendas", "card", 330, 22, 290, 125, {"Fields": [field_measure("Vendas")]}, "Vendas"), visual("Card_Unidades", "card", 632, 22, 290, 125, {"Fields": [field_measure("Unidades")]}, "Unidades vendidas"), visual("Card_Ticket", "card", 934, 22, 318, 125, {"Fields": [field_measure("Ticket médio")]}, "Ticket médio")]
    page1_visuals = cards + [
        visual("Trend_Faturamento", "lineChart", 28, 165, 600, 300, {"Category": [field_column("Data")], "Y": [field_measure("Faturamento")]}, "Evolução do faturamento"),
        visual("Rank_Vendedores", "clusteredBarChart", 646, 165, 606, 300, {"Category": [field_column("Vendedor")], "Y": [field_measure("Faturamento"), field_measure("Vendas")]}, "Desempenho por vendedor"),
        visual("Rank_Produtos", "clusteredBarChart", 28, 482, 600, 300, {"Category": [field_column("Produto")], "Y": [field_measure("Unidades"), field_measure("Faturamento")]}, "Produtos: unidades e faturamento"),
        visual("Tabela_Vendas", "tableEx", 646, 482, 606, 300, {"Values": [field_column("Data"), field_column("Vendedor"), field_column("Produto"), field_column("Quantidade"), field_column("Valor unitário")]}, "Detalhamento das vendas"),
    ]
    page2_visuals = [
        visual("Seller_Table", "tableEx", 28, 24, 600, 650, {"Values": [field_column("Vendedor"), field_measure("Vendas"), field_measure("Unidades"), field_measure("Faturamento"), field_measure("Ticket médio")]}, "Resultado por vendedor"),
        visual("Product_Table", "tableEx", 646, 24, 606, 650, {"Values": [field_column("Produto"), field_measure("Unidades"), field_measure("Vendas"), field_measure("Faturamento")]}, "Resultado por produto"),
    ]
    def json_bytes(obj: Any) -> bytes:
        return json.dumps(obj, ensure_ascii=False, indent=2).encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{project}/{project}.pbip", json_bytes(pbip))
        archive.writestr(f"{project}/.gitignore", "**/.pbi/localSettings.json\n**/.pbi/cache.abf\n")
        archive.writestr(f"{project}/{model_name}/definition.pbism", json_bytes({"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/semanticModel/definitionProperties/1.0.0/schema.json", "version": "1.0"}))
        archive.writestr(f"{project}/{model_name}/model.bim", model_json.encode("utf-8"))
        archive.writestr(f"{project}/{report_name}/definition.pbir", json_bytes(pbir))
        archive.writestr(f"{project}/{report_name}/definition/version.json", json_bytes({"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/versionMetadata/1.0.0/schema.json", "version": "2.0.0"}))
        theme_file = "SalesFlowTheme-7a31c9e2.json"
        theme = {
            "name": theme_file,
            "dataColors": ["#2563EB", "#128B83", "#7C3AED", "#F59E0B", "#14B8A6", "#D85C5C", "#4F75C2", "#9A6DD7"],
            "background": "#F4F7FB",
            "foreground": "#172B46",
            "firstLevelElements": "#172B46",
            "secondLevelElements": "#52657D",
            "thirdLevelElements": "#EDF3FF",
            "fourthLevelElements": "#E4EAF1",
            "tableAccent": "#2563EB",
            "good": "#128B83",
            "neutral": "#F59E0B",
            "bad": "#D85C5C",
            "textClasses": {
                "title": {"fontFace": "Segoe UI Semibold", "fontSize": 13, "color": "#172B46"},
                "callout": {"fontFace": "Segoe UI Semibold", "fontSize": 28, "color": "#172B46"},
                "label": {"fontFace": "Segoe UI", "fontSize": 10, "color": "#738298"},
            },
            "visualStyles": {
                "tableEx": {
                    "*": {
                        "columnHeaders": [{"fontColor": {"solid": {"color": "#52657D"}}, "backColor": {"solid": {"color": "#EDF3FF"}}, "wordWrap": False}],
                        "values": [{"fontColorPrimary": {"solid": {"color": "#334963"}}, "backColorPrimary": {"solid": {"color": "#FFFFFF"}}, "fontColorSecondary": {"solid": {"color": "#334963"}}, "backColorSecondary": {"solid": {"color": "#F7F9FC"}}}],
                        "grid": [{"outlineWeight": 1}],
                    }
                }
            },
        }
        theme_metadata = {"name": theme_file, "reportVersionAtImport": {"visual": "2.7.0", "report": "3.3.0", "page": "2.1.0"}, "type": "RegisteredResources"}
        report = {
            "$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/report/3.3.0/schema.json",
            "themeCollection": {
                "baseTheme": {"name": "CY24SU06", "reportVersionAtImport": {"visual": "2.7.0", "report": "3.3.0", "page": "2.1.0"}, "type": "SharedResources"},
                "customTheme": theme_metadata,
            },
            "resourcePackages": [{"name": "RegisteredResources", "type": "RegisteredResources", "items": [{"name": theme_file, "path": theme_file, "type": "CustomTheme"}]}],
            "filterConfig": {"filters": []},
        }
        archive.writestr(f"{project}/{report_name}/definition/report.json", json_bytes(report))
        archive.writestr(f"{project}/{report_name}/StaticResources/RegisteredResources/{theme_file}", json_bytes(theme))
        page_ids = {page_name: "ReportSection" + hashlib.sha1(f"{project}:{page_name}".encode("utf-8")).hexdigest()[:20] for page_name in ("VisaoGeral", "AnaliseDetalhada")}
        archive.writestr(f"{project}/{report_name}/definition/pages/pages.json", json_bytes({"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/pagesMetadata/1.0.0/schema.json", "pageOrder": [page_ids["VisaoGeral"], page_ids["AnaliseDetalhada"]], "activePageName": page_ids["VisaoGeral"]}))
        for page_name, display_name, visuals in [("VisaoGeral", "Visão geral", page1_visuals), ("AnaliseDetalhada", "Análise detalhada", page2_visuals)]:
            folder = f"{project}/{report_name}/definition/pages/{page_ids[page_name]}"
            page = {"$schema": "https://developer.microsoft.com/json-schemas/fabric/item/report/definition/page/2.1.0/schema.json", "name": page_ids[page_name], "displayName": display_name, "displayOption": "FitToPage", "height": 900, "width": 1280, "filterConfig": {"filters": []}, "objects": {"background": [{"properties": {"color": fill("#F4F7FB"), "transparency": literal("0D")}}], "outspace": [{"properties": {"color": fill("#F4F7FB"), "transparency": literal("0D")}}]}}
            archive.writestr(f"{folder}/page.json", json_bytes(page))
            for item in visuals:
                visual_id = hashlib.sha1(f"{page_name}:{item['name']}".encode("utf-8")).hexdigest()[:20]
                visual_item = {**item, "name": visual_id}
                archive.writestr(f"{folder}/visuals/{visual_id}/visual.json", json_bytes(visual_item))
        archive.writestr(f"{project}/LEIA-ME.txt", "Projeto PBIP do SalesFlow. Extraia toda a pasta e abra SalesFlow_Vendas.pbip no Power BI Desktop. Depois use Arquivo > Salvar como > Power BI Desktop (.pbix).\n\nO projeto contém duas páginas: Visão geral (faturamento, vendas, unidades, ticket médio, tendência, ranking de vendedores/produtos e detalhamento) e Análise detalhada (tabelas por vendedor e produto). Os filtros de Data, Vendedor e Produto podem ser adicionados no Desktop como segmentações para navegação interativa.\n")
    return out.getvalue()


@app.post("/api/runs/{run_id}/powerbi-export")
def export_run_powerbi(run_id: int, payload: dict[str, Any] | None = None):
    payload = payload or {}
    with Session() as db:
        run = db.get(Run, run_id)
        if not run:
            raise HTTPException(404, "Execução não encontrada.")
        corrections = corrections_for_export(run, payload)
    rows, omitted_rows, _ = dashboard_rows_for_run(run, corrections)
    if not rows:
        raise HTTPException(422, "Esta execução não contém vendas válidas para montar o projeto Power BI.")
    archive_bytes = powerbi_project(rows, run_id)
    filename = f"SalesFlow_Vendas_{run_id}.zip"
    return StreamingResponse(
        io.BytesIO(archive_bytes),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "X-Omitted-Rows": str(omitted_rows)},
    )


@app.get("/api/runs/{run_id}")
def get_run(run_id: int):
    with Session() as db:
        run = db.get(Run, run_id)
        if not run:
            raise HTTPException(404, "Execução não encontrada.")
        return run_dict(run)


@app.get("/api/reports/export")
def export(format: str = "json"):
    if format == "json":
        return JSONResponse(summary())
    if format == "csv":
        with Session() as db:
            sales = db.scalars(select(Sale).order_by(Sale.sale_date)).all()
            rows = [{"data": s.sale_date.isoformat(), "vendedor": s.seller, "produto": s.product, "quantidade": s.quantity, "valor_unitario": s.unit_price, "faturamento": round(s.quantity * s.unit_price, 2)} for s in sales]
        csv = pd.DataFrame(rows).to_csv(index=False, encoding="utf-8-sig")
        return StreamingResponse(io.StringIO(csv), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename= vendas_consolidadas.csv"})
    raise HTTPException(400, "Formato inválido. Use json ou csv.")
